"""Apply verified batches or refresh downloaded public datasets, preserving revisions."""
from pathlib import Path
import argparse,copy,csv,datetime,json,os
from openpyxl import load_workbook
from build import OUT,load,write_json,validate,build
def merge(d,key,points):
 old={p['period']:p for p in d['series'][key]['points']}; added=revised=0;seen=set()
 for p in points:
  assert p['period'] not in seen,(key,'duplicate incoming period');seen.add(p['period'])
  if p['period'] not in old:added+=1
  elif p!=old[p['period']]:revised+=1
  old[p['period']]=p
 d['series'][key]['points']=sorted(old.values(),key=lambda p:p['period'])
 return added,revised
def run():
 ap=argparse.ArgumentParser();ap.add_argument('--incoming',type=Path);ap.add_argument('--batch',type=Path);ap.add_argument('--as-of',default=datetime.date.today().isoformat())
 a=ap.parse_args();old=load(OUT/'data.json');d=copy.deepcopy(old);d['updated_at']=a.as_of;events=[];changes=0
 if a.incoming:
  manifest=load(a.incoming/'manifest.json')
  for item in manifest:
   event={'source':item['id'],'status':item['status'],'as_of':a.as_of}
   if item['status']!='downloaded':event['error']=item.get('error','下载失败，沿用旧值');events.append(event);continue
   trial=copy.deepcopy(d);count=0
   try:
    file=a.incoming/item['file'];key=item['id']
    if key=='greix':
     raw=load(file)['chart_legend']['3'];buckets={x:[] for x in ['berlin_price','munich_price','frankfurt_price']}
     for period,row in raw.items():
      q,y=period.split()
      for en in ['Berlin','Munich','Frankfurt']:
       if en in row and row[en]['data_value'] is not None:
        p=row[en];buckets[en.lower()+'_price'].append({'period':f'{y}-{q}','value':p['data_value'],'source_id':'greix','official_yoy':p.get('data_growth_value')})
    elif key=='ecb':
     buckets={'de_rate':[{'period':r['TIME_PERIOD'],'value':float(r['OBS_VALUE']),'source_id':'ecb'} for r in csv.DictReader(file.open(encoding='utf-8-sig'))]}
    elif key=='pmms':
     pts=[]
     for row in load_workbook(file,read_only=True,data_only=True).active.iter_rows(values_only=True):
      if isinstance(row[0],datetime.datetime) and row[0].year>=2015 and isinstance(row[1],(int,float)):
       pts.append({'period':row[0].strftime('%Y-%m-%d'),'value':row[1],'source_id':'pmms'})
     buckets={'us_rate':pts}
    else:raise ValueError('Unknown source '+key)
    details=[]
    for sk,pts in buckets.items():
     assert len(pts)>0,'empty response'
     prior=trial['series'][sk]['points'];old_end=prior[-1]['period']
     pts.sort(key=lambda p:p['period'])
     assert pts[-1]['period']>=old_end,'source returned an older cutoff'
     absent={p['period'] for p in prior}-{p['period'] for p in pts}
     n,r=merge(trial,sk,pts);count+=n+r
     details.append({'series':sk,'added':n,'revised':r,'missing_retained':sorted(absent)})
    validate(trial);trial['sources'][key]['accessed_at']=a.as_of
    d=trial;changes+=count;event.update(status='checked',changes=count,details=details)
   except Exception as e:event.update(status='failed',error=str(e))
   events.append(event)
 if a.batch:
  b=load(a.batch)
  assert set(b)<= {'sources','series','analysis'},'unknown batch fields'
  if any(d['sources'].get(k)!=v for k,v in b.get('sources',{}).items()):changes+=1
  d['sources'].update(b.get('sources',{}))
  for key,s in b.get('series',{}).items():
   assert key in d['series'],'new series requires explicit scope and schema review'
   assert set(s)=={'points'},'batch may only update points of an existing series'
   n,r=merge(d,key,s['points']);changes+=n+r
  if 'analysis' in b:d['analysis']=b['analysis'];changes+=1
  # A revised Beijing monthly ratio changes the entire subsequent chain.
  total=100.
  for i,p in enumerate(d['series']['beijing_price']['points']):
   if i:total*=p['mom_index']/100
   p['value']=total
  validate(d);events.append({'source':'verified_manual_batch','status':'checked','as_of':a.as_of,'file':a.batch.name})
 if changes:
  validate(d)
  history=OUT/'history';history.mkdir(exist_ok=True)
  stamp=datetime.datetime.now().strftime('%Y%m%d-%H%M%S-%f')
  write_json(history/(stamp+'.json'),old)
  write_json(OUT/'data.pending.json',d);os.replace(OUT/'data.pending.json',OUT/'data.json')
  for item in (load(a.incoming/'manifest.json') if a.incoming else []):
   if any(e['source']==item['id'] and e['status']=='checked' for e in events):
    import shutil
    shutil.copy2(a.incoming/item['file'],OUT/'source_data'/item['file'])
 with (OUT/'refresh_log.jsonl').open('a',encoding='utf8') as f:
  f.write(json.dumps({'checked_at':a.as_of,'changed_points':changes,'events':events},ensure_ascii=False)+'\n')
 build()
 print(json.dumps({'changed_points':changes,'sources':events},ensure_ascii=False))
 if any(e['status'] in ('failed','download_failed') for e in events):raise SystemExit(2)
if __name__=='__main__':run()
